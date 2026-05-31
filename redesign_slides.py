"""
Redesign cbc-auction-2026.pptx:
  - Replace the dark navy colour palette with a vibrant blue + orange/gold scheme
  - Increase all font sizes for maximum legibility
  - Preserve all layout, positions, text content and player data

Usage:
    python3 redesign_slides.py [--src ./cbc-auction-2026.pptx] [--dest ./cbc-auction-2026.pptx]
"""
import argparse
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

# ── Colour maps ───────────────────────────────────────────────────────────────

# Old fill colour → new fill colour  (hex strings, upper-case, no '#')
FILL_MAP = {
    # Dark navy backgrounds → vibrant blue
    '040D1E': '1565C0',
    '071428': '0D47A1',
    '003A55': '0D47A1',
    '001414': '0D47A1',
    '0D203A': '1565C0',
    '141400': '1565C0',
    '141414': '1565C0',
    # Detail-card dark blues → rich dark indigo (keeps white/gold text readable)
    '0A1A30': '0D2A6B',
    '081426': '0D2A6B',
    '0B1E3F': '0D2A6B',
    # Generic dark separator
    '2A3A4A': '1976D2',
    # Cyan accent → vivid amber-orange
    '00D4FF': 'FF8F00',
    # Round badge dark → vivid orange
    '002A40': 'FF8F00',
    # Category / team badge colours — keep distinctive
    '00CC77': '00CC77',
    '00FFAA': '69F0AE',
    'BB44FF': 'CE93D8',
    'FF7700': 'FF7700',
    'FFD700': 'FFD700',
    'FFE600': 'FFE600',
    'FF3333': 'FF3333',
}

# Old text (font) colour → new text colour
TEXT_MAP = {
    'FFFFFF': 'FFFFFF',   # white stays — now on blue backgrounds
    '040D1E': '040D1E',   # dark text stays (footer on bright bars)
    '00D4FF': 'FFD54F',   # cyan text → warm amber/gold
    '7896AB': 'B0BEC5',   # muted blue-grey → lighter, visible on blue bg
    'CCDDEE': 'CFD8DC',   # very-light text → stays light
    'FFE600': 'FFE600',   # gold price text stays
    'FF7700': 'FF7700',   # orange text stays
    'FFD700': 'FFD700',   # gold text stays
    '00FFAA': '69F0AE',   # green text — brighter
    'BB44FF': 'CE93D8',   # purple text — lighter
    '00CC77': '69F0AE',   # team-assigned green
}

# ── Special per-shape overrides ───────────────────────────────────────────────
# shape.name → { 'fill': 'RRGGBB', 'text': 'RRGGBB' }
SHAPE_OVERRIDES = {
    # Round badge: vivid orange fill + white text for maximum pop
    'round_badge': {'fill': 'E65100', 'text': 'FFFFFF'},
    # Logo oval: amber fill + dark text
    'logo':        {'fill': 'FF8F00', 'text': '040D1E'},
    'Oval 9':      {'fill': 'FF8F00', 'text': '040D1E'},
    'Oval 11':     {'fill': 'FF8F00', 'text': '040D1E'},
}

# ── Font-size bumps ───────────────────────────────────────────────────────────
def bump_size(pt: float) -> float:
    """Return a bigger font size — stepped increases to avoid overflow."""
    if pt <= 10:
        return pt + 2
    if pt <= 14:
        return pt + 3
    if pt <= 20:
        return pt + 4
    if pt <= 30:
        return pt + 5
    if pt <= 50:
        return pt + 7
    return pt + 10   # very large titles


# ── Helpers ───────────────────────────────────────────────────────────────────
def rgb(hex_str: str) -> RGBColor:
    h = hex_str.lstrip('#')
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def apply_fill(shape, new_hex: str):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(new_hex)


def remap_fill(shape):
    try:
        fill = shape.fill
        if str(fill.type) != 'SOLID (1)':
            return
        old = str(fill.fore_color.rgb).upper()
        new = FILL_MAP.get(old)
        if new:
            apply_fill(shape, new)
    except Exception:
        pass


def remap_text(shape, text_override: str | None = None):
    if not shape.has_text_frame:
        return
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            # Font size
            if run.font.size:
                old_pt = run.font.size.pt
                run.font.size = Pt(bump_size(old_pt))
            # Font colour
            if text_override:
                run.font.color.rgb = rgb(text_override)
            elif run.font.color and run.font.color.type:
                try:
                    old = str(run.font.color.rgb).upper()
                    new = TEXT_MAP.get(old)
                    if new:
                        run.font.color.rgb = rgb(new)
                except Exception:
                    pass


# ── Main ──────────────────────────────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--src',  default='./cbc-auction-2026.pptx')
    parser.add_argument('--dest', default='./cbc-auction-2026.pptx')
    args = parser.parse_args()

    print(f'Opening {args.src} …')
    prs = Presentation(args.src)

    for slide_idx, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            override = SHAPE_OVERRIDES.get(shape.name, {})

            # Fill
            fill_hex = override.get('fill')
            if fill_hex:
                try:
                    apply_fill(shape, fill_hex)
                except Exception:
                    pass
            else:
                remap_fill(shape)

            # Text / font
            text_hex = override.get('text')
            remap_text(shape, text_override=text_hex)

    prs.save(args.dest)
    print(f'Saved redesigned presentation → {args.dest}')
    print(f'Total slides processed: {len(prs.slides)}')


if __name__ == '__main__':
    main()
